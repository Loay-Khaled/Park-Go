import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette Definitions
    COLOR_BG = RGBColor(11, 19, 43)          # #0B132B Deep Executive Navy
    COLOR_BG_CARD = RGBColor(22, 32, 54)     # #162036 Surface Card Navy
    COLOR_BORDER = RGBColor(37, 99, 235)     # #2563EB Electric Blue
    COLOR_BORDER_MUTED = RGBColor(30, 41, 59) # #1E293B Subtle Slate Border
    COLOR_WHITE = RGBColor(255, 255, 255)    # #FFFFFF Pure White
    COLOR_MUTED = RGBColor(148, 163, 184)    # #94A3B8 Slate Grey
    COLOR_TEXT_BODY = RGBColor(203, 213, 225) # #CBD5E1 Light Slate Body
    COLOR_ACCENT = RGBColor(59, 130, 246)    # #3B82F6 Vibrant Blue
    COLOR_CYAN = RGBColor(6, 182, 212)       # #06B6D4 Electric Cyan
    COLOR_GREEN = RGBColor(16, 185, 129)     # #10B981 Emerald Accent

    def set_slide_background(slide):
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5)
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = COLOR_BG
        bg_shape.line.fill.background()

        # Add decorative subtle top glow bar
        glow_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), 0, Inches(4.0), Inches(0.06)
        )
        glow_bar.fill.solid()
        glow_bar.fill.fore_color.rgb = COLOR_ACCENT
        glow_bar.line.fill.background()

    def add_slide_header(slide, tracker, title, subtitle):
        # Tracker Badge
        tracker_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.6), Inches(3.4), Inches(0.36)
        )
        tracker_box.fill.solid()
        tracker_box.fill.fore_color.rgb = COLOR_BG_CARD
        tracker_box.line.color.rgb = COLOR_BORDER
        tracker_box.line.width = Pt(1)
        tf_tr = tracker_box.text_frame
        tf_tr.word_wrap = True
        tf_tr.margin_left = Inches(0.15)
        tf_tr.margin_right = Inches(0.15)
        tf_tr.margin_top = Inches(0.05)
        tf_tr.margin_bottom = Inches(0.05)
        p_tr = tf_tr.paragraphs[0]
        p_tr.text = tracker.upper()
        p_tr.font.size = Pt(9.5)
        p_tr.font.bold = True
        p_tr.font.color.rgb = COLOR_CYAN
        p_tr.font.name = "Calibri"

        # Main Title & Subtitle Box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.05), Inches(11.733), Inches(1.15))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(26)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_WHITE
        p1.font.name = "Calibri"

        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(13.5)
        p2.font.color.rgb = COLOR_MUTED
        p2.font.name = "Calibri"

    def add_footer(slide, current_slide, total_slides=6):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.85), Inches(11.733), Inches(0.4))
        tf = footer_box.text_frame
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = f"Park & Go  |  Strategic Partnership Proposal — Smart Village, Egypt               Slide {current_slide} of {total_slides}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_MUTED
        p.font.name = "Calibri"

    # =========================================================================
    # SLIDE 1: Title & Introduction
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Decorative geometric accent on right
    accent_box = s1.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(1.2), Inches(3.7), Inches(5.0)
    )
    accent_box.fill.solid()
    accent_box.fill.fore_color.rgb = COLOR_BG_CARD
    accent_box.line.color.rgb = COLOR_BORDER
    accent_box.line.width = Pt(1.5)
    tf_acc = accent_box.text_frame
    tf_acc.word_wrap = True
    tf_acc.margin_left = Inches(0.4)
    tf_acc.margin_right = Inches(0.4)
    tf_acc.margin_top = Inches(0.5)

    p_a1 = tf_acc.paragraphs[0]
    p_a1.text = "PROPOSAL SUMMARY"
    p_a1.font.size = Pt(11)
    p_a1.font.bold = True
    p_a1.font.color.rgb = COLOR_CYAN

    p_a2 = tf_acc.add_paragraph()
    p_a2.text = "\n• Audience:\nSmart Village Management, Egypt\n\n• Focus:\nCorporate Mobility & Operational Optimization\n\n• Core Model:\nComplementary to Existing Leased Parking Framework\n\n• Engagement:\nExploratory Strategic Collaboration & Pilot Initiative"
    p_a2.font.size = Pt(12)
    p_a2.font.color.rgb = COLOR_TEXT_BODY

    # Left content block
    # Strategic Badge
    s1_badge = s1.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(4.5), Inches(0.42)
    )
    s1_badge.fill.solid()
    s1_badge.fill.fore_color.rgb = COLOR_BG_CARD
    s1_badge.line.color.rgb = COLOR_BORDER
    s1_badge.line.width = Pt(1)
    tf_b = s1_badge.text_frame
    tf_b.margin_left = Inches(0.2)
    tf_b.margin_top = Inches(0.08)
    p_b = tf_b.paragraphs[0]
    p_b.text = "STRATEGIC PARTNERSHIP PROPOSAL"
    p_b.font.size = Pt(10.5)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_CYAN

    # Company & Tagline
    title_box_s1 = s1.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(7.6), Inches(2.6))
    tf_s1 = title_box_s1.text_frame
    tf_s1.word_wrap = True
    tf_s1.margin_left = 0
    tf_s1.margin_top = 0

    p_main = tf_s1.paragraphs[0]
    p_main.text = "Park & Go"
    p_main.font.size = Pt(46)
    p_main.font.bold = True
    p_main.font.color.rgb = COLOR_WHITE
    p_main.font.name = "Calibri"

    p_tag = tf_s1.add_paragraph()
    p_tag.text = "Smart Parking, Smarter Opportunities."
    p_tag.font.size = Pt(22)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_ACCENT
    p_tag.font.name = "Calibri"

    p_desc = tf_s1.add_paragraph()
    p_desc.text = "\nIntroducing a smart digital mobility platform tailored for Egypt's leading business communities—engineered to elevate parking convenience, enhance operational visibility, and unlock new value from existing infrastructure."
    p_desc.font.size = Pt(13)
    p_desc.font.color.rgb = COLOR_TEXT_BODY

    # Meta card (Prepared For / Presented By)
    meta_box = s1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(4.85), Inches(7.6), Inches(1.35)
    )
    meta_box.fill.solid()
    meta_box.fill.fore_color.rgb = COLOR_BG_CARD
    meta_box.line.color.rgb = COLOR_BORDER_MUTED
    tf_meta = meta_box.text_frame
    tf_meta.word_wrap = True
    tf_meta.margin_left = Inches(0.3)
    tf_meta.margin_top = Inches(0.2)

    p_m1 = tf_meta.paragraphs[0]
    p_m1.text = "PREPARED FOR: Smart Village Management, Giza, Egypt"
    p_m1.font.size = Pt(11.5)
    p_m1.font.bold = True
    p_m1.font.color.rgb = COLOR_WHITE

    p_m2 = tf_meta.add_paragraph()
    p_m2.text = "PRESENTED BY: Loay Khaled — Founder, Park & Go"
    p_m2.font.size = Pt(11.5)
    p_m2.font.bold = True
    p_m2.font.color.rgb = COLOR_ACCENT

    p_m3 = tf_meta.add_paragraph()
    p_m3.text = "PURPOSE: Strategic Exploration & Mutually Beneficial Partnership"
    p_m3.font.size = Pt(10.5)
    p_m3.font.color.rgb = COLOR_MUTED

    add_footer(s1, 1)

    # =========================================================================
    # SLIDE 2: The Opportunity
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_slide_header(
        s2,
        "02. The Market Context",
        "The Corporate Mobility Opportunity",
        "Unlocking greater efficiency, convenience, and visibility from high-density corporate parking assets."
    )

    # 3 Horizontal Cards
    card_width = Inches(3.68)
    card_height = Inches(4.1)
    card_y = Inches(2.35)
    card_gap = Inches(0.34)

    s2_data = [
        {
            "num": "01",
            "title": "Optimized Space Utilization",
            "tag": "Capacity Optimization",
            "points": [
                "Corporate campuses experience dynamic occupancy surges between tenant staff, commercial guests, and event visitors.",
                "Digital management bridges the gap between peak morning arrivals and mid-day meeting turnovers.",
                "Maximizes the productive utility of available bays without altering allocated lease spaces."
            ]
        },
        {
            "num": "02",
            "title": "Elevated Tenant & Visitor Experience",
            "tag": "Executive Convenience",
            "points": [
                "Parking is the primary physical touchpoint for every C-level executive, employee, and multinational client entering the campus.",
                "Frictionless digital navigation and coordination save valuable business hours previously lost to circling.",
                "Directly reinforces Smart Village's reputation as a world-class, premium business destination."
            ]
        },
        {
            "num": "03",
            "title": "Operational Visibility & Data Insight",
            "tag": "Campus Intelligence",
            "points": [
                "Traditional parking models often operate with limited real-time data on actual utilization patterns and peak flow hours.",
                "Smart digital coordination delivers comprehensive operational transparency across all facility zones.",
                "Empowers leadership with empirical insights for smarter long-term master planning and logistics."
            ]
        }
    ]

    for i, c in enumerate(s2_data):
        x = Inches(0.8) + i * (card_width + card_gap)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, card_y, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_BG_CARD
        card.line.color.rgb = COLOR_BORDER if i == 1 else COLOR_BORDER_MUTED
        card.line.width = Pt(1.5 if i == 1 else 1)

        # Top Accent Ribbon
        ribbon = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Inches(0.3), card_y, Inches(1.0), Inches(0.06))
        ribbon.fill.solid()
        ribbon.fill.fore_color.rgb = COLOR_ACCENT
        ribbon.line.fill.background()

        tf_c = card.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.32)
        tf_c.margin_right = Inches(0.32)
        tf_c.margin_top = Inches(0.3)

        p_num = tf_c.paragraphs[0]
        p_num.text = f"PILLAR {c['num']}  •  {c['tag'].upper()}"
        p_num.font.size = Pt(9)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_CYAN

        p_t = tf_c.add_paragraph()
        p_t.text = c["title"]
        p_t.font.size = Pt(16.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE

        for pt in c["points"]:
            p_pt = tf_c.add_paragraph()
            p_pt.text = f"•  {pt}"
            p_pt.font.size = Pt(11.5)
            p_pt.font.color.rgb = COLOR_TEXT_BODY

    add_footer(s2, 2)

    # =========================================================================
    # SLIDE 3: Our Concept
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_slide_header(
        s3,
        "03. Solution Overview",
        "The Park & Go Concept",
        "A sophisticated digital layer designed to seamlessly coordinate parking across high-profile commercial ecosystems."
    )

    # Top Concept Banner
    hero_card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.25), Inches(11.733), Inches(1.15))
    hero_card.fill.solid()
    hero_card.fill.fore_color.rgb = COLOR_BG_CARD
    hero_card.line.color.rgb = COLOR_BORDER
    hero_card.line.width = Pt(1.5)
    tf_h = hero_card.text_frame
    tf_h.word_wrap = True
    tf_h.margin_left = Inches(0.35)
    tf_h.margin_right = Inches(0.35)
    tf_h.margin_top = Inches(0.18)

    p_h1 = tf_h.paragraphs[0]
    p_h1.text = "CORE VALUE PROPOSITION"
    p_h1.font.size = Pt(9.5)
    p_h1.font.bold = True
    p_h1.font.color.rgb = COLOR_CYAN

    p_h2 = tf_h.add_paragraph()
    p_h2.text = "Park & Go serves as an intelligent software-driven layer that bridges property managers, corporate tenants, and visiting motorists—transforming static parking infrastructure into a coordinated, customer-centric digital amenity."
    p_h2.font.size = Pt(12.5)
    p_h2.font.color.rgb = COLOR_WHITE

    # 3 Solution Focus Areas
    col_w = Inches(3.68)
    col_h = Inches(3.0)
    col_y = Inches(3.6)

    s3_pillars = [
        {
            "title": "Frictionless Corporate Access",
            "badge": "FOR EMPLOYEES & VISITORS",
            "desc": "Enables pre-arrival coordination and streamlined campus entry. Eliminates manual verification delays, ensuring that executive guests and workforce arrive calmly and on schedule."
        },
        {
            "title": "Intelligent Space Orchestration",
            "badge": "FOR CAMPUS OPERATIONS",
            "desc": "Dynamically guides incoming vehicular flow to available zones without operational friction. Balances campus density across multi-level and surface lots effortlessly."
        },
        {
            "title": "Executive Intelligence Dashboard",
            "badge": "FOR SMART VILLAGE MANAGEMENT",
            "desc": "Provides high-level, aggregate operational reporting on capacity utilization, peak transit windows, and facility dynamics to support strategic asset oversight."
        }
    ]

    for i, p in enumerate(s3_pillars):
        x = Inches(0.8) + i * (col_w + card_gap)
        c_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, col_y, col_w, col_h)
        c_box.fill.solid()
        c_box.fill.fore_color.rgb = COLOR_BG_CARD
        c_box.line.color.rgb = COLOR_BORDER_MUTED
        c_box.line.width = Pt(1)

        tf_p = c_box.text_frame
        tf_p.word_wrap = True
        tf_p.margin_left = Inches(0.3)
        tf_p.margin_right = Inches(0.3)
        tf_p.margin_top = Inches(0.25)

        p_b = tf_p.paragraphs[0]
        p_b.text = p["badge"]
        p_b.font.size = Pt(9)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_ACCENT

        p_t = tf_p.add_paragraph()
        p_t.text = p["title"]
        p_t.font.size = Pt(15.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE

        p_d = tf_p.add_paragraph()
        p_d.text = f"\n{p['desc']}"
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = COLOR_TEXT_BODY

    add_footer(s3, 3)

    # =========================================================================
    # SLIDE 4: Why Smart Village?
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_slide_header(
        s4,
        "04. Strategic Alignment",
        "Why Smart Village? A Fully Complementary Approach",
        "Engineered to enhance campus operations without changing tenants' allocated parking spaces."
    )

    # Core Reassurance Highlight Banner
    reassure_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.25), Inches(11.733), Inches(0.85))
    reassure_box.fill.solid()
    reassure_box.fill.fore_color.rgb = COLOR_BG_CARD
    reassure_box.line.color.rgb = COLOR_GREEN
    reassure_box.line.width = Pt(1.5)
    tf_re = reassure_box.text_frame
    tf_re.word_wrap = True
    tf_re.margin_left = Inches(0.35)
    tf_re.margin_top = Inches(0.18)

    p_r1 = tf_re.paragraphs[0]
    p_r1.text = "FOUNDATIONAL PRINCIPLE: 100% PRESERVATION OF TENANT ALLOCATIONS"
    p_r1.font.size = Pt(9.5)
    p_r1.font.bold = True
    p_r1.font.color.rgb = COLOR_GREEN

    p_r2 = tf_re.add_paragraph()
    p_r2.text = "Our proposed model operates strictly alongside Smart Village's existing commercial leasing agreements. It does NOT reallocate, reduce, or alter contractual parking spots reserved for corporate tenants."
    p_r2.font.size = Pt(11.5)
    p_r2.font.color.rgb = COLOR_WHITE

    # 3 Alignment Pillars
    s4_items = [
        {
            "tag": "CONTRACTUAL INTEGRITY",
            "title": "Complements Existing Rental Models",
            "desc": "Smart Village's core parking revenue originates from corporate long-term leases. Park & Go focuses exclusively on optimizing peripheral, visitor, or unallocated capacity—introducing new value without encroaching on established commitments."
        },
        {
            "tag": "TECH LEADERSHIP",
            "title": "Matches Smart Village's DNA",
            "desc": "As Egypt's premier hub for technology and multinational enterprises, Smart Village is the natural home for forward-looking smart mobility solutions. Adopting digital parking aligns directly with the campus's modern identity."
        },
        {
            "tag": "ZERO FRICTION",
            "title": "Non-Intrusive, Phased Rollout",
            "desc": "Designed to integrate smoothly with existing security standards and entry protocols. Implementation can begin on a targeted, low-profile scale to ensure complete management comfort before broader deployment."
        }
    ]

    for i, it in enumerate(s4_items):
        x = Inches(0.8) + i * (col_w + card_gap)
        box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(3.3), col_w, Inches(3.3))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_BG_CARD
        box.line.color.rgb = COLOR_BORDER_MUTED
        box.line.width = Pt(1)

        tf_it = box.text_frame
        tf_it.word_wrap = True
        tf_it.margin_left = Inches(0.3)
        tf_it.margin_right = Inches(0.3)
        tf_it.margin_top = Inches(0.25)

        p_tag = tf_it.paragraphs[0]
        p_tag.text = it["tag"]
        p_tag.font.size = Pt(9)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_CYAN

        p_t = tf_it.add_paragraph()
        p_t.text = it["title"]
        p_t.font.size = Pt(15.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE

        p_d = tf_it.add_paragraph()
        p_d.text = f"\n{it['desc']}"
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = COLOR_TEXT_BODY

    add_footer(s4, 4)

    # =========================================================================
    # SLIDE 5: Partnership Benefits
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_slide_header(
        s5,
        "05. Mutual Value Creation",
        "Partnership Benefits for Smart Village",
        "A strategic collaboration structured to deliver commercial, operational, and reputational upside."
    )

    # 4 Cards in 2x2 Grid
    grid_w = Inches(5.7)
    grid_h = Inches(2.05)
    grid_gap_x = Inches(0.33)
    grid_gap_y = Inches(0.25)
    grid_y = Inches(2.3)

    benefits = [
        {
            "num": "01",
            "title": "Potential Incremental Recurring Revenue",
            "tag": "COMMERCIAL VALUE",
            "desc": "Opportunity to unlock additional recurring returns from commercial visitor parking and event overflow, monetizing spaces during off-peak hours without operational capital expense."
        },
        {
            "num": "02",
            "title": "Premium Digital Amenities for Tenants",
            "tag": "SERVICE ENHANCEMENT",
            "desc": "Equip tenant corporations with modern parking coordination capabilities, elevating tenant satisfaction, boosting lease renewal retention, and attracting prospective multinational tenants."
        },
        {
            "num": "03",
            "title": "Reduced Campus Congestion & Friction",
            "tag": "OPERATIONAL EFFICIENCY",
            "desc": "Mitigate morning queue backups at main access security gates through coordinated arrivals, resulting in smoother campus traffic flow and less strain on on-site security personnel."
        },
        {
            "num": "04",
            "title": "Environmental & Sustainability Impact",
            "tag": "ESG & SUSTAINABILITY",
            "desc": "Significantly decreases intra-campus vehicle idling and carbon emissions, directly contributing to Smart Village's environmental sustainability benchmarks and corporate green image."
        }
    ]

    for i, b in enumerate(benefits):
        col = i % 2
        row = i // 2
        x = Inches(0.8) + col * (grid_w + grid_gap_x)
        y = grid_y + row * (grid_h + grid_gap_y)

        b_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, grid_w, grid_h)
        b_card.fill.solid()
        b_card.fill.fore_color.rgb = COLOR_BG_CARD
        b_card.line.color.rgb = COLOR_BORDER_MUTED
        b_card.line.width = Pt(1)

        tf_b = b_card.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = Inches(0.35)
        tf_b.margin_right = Inches(0.35)
        tf_b.margin_top = Inches(0.2)

        p_t1 = tf_b.paragraphs[0]
        p_t1.text = f"BENEFIT {b['num']}  •  {b['tag']}"
        p_t1.font.size = Pt(8.5)
        p_t1.font.bold = True
        p_t1.font.color.rgb = COLOR_CYAN

        p_t2 = tf_b.add_paragraph()
        p_t2.text = b["title"]
        p_t2.font.size = Pt(15)
        p_t2.font.bold = True
        p_t2.font.color.rgb = COLOR_WHITE

        p_t3 = tf_b.add_paragraph()
        p_t3.text = b["desc"]
        p_t3.font.size = Pt(11)
        p_t3.font.color.rgb = COLOR_TEXT_BODY

    add_footer(s5, 5)

    # =========================================================================
    # SLIDE 6: Let's Collaborate
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_slide_header(
        s6,
        "06. Next Steps",
        "Let's Collaborate",
        "A structured, low-risk path forward to evaluate mutual synergies."
    )

    # 3 Next Step Action Cards
    s6_steps = [
        {
            "step": "STEP 1",
            "title": "Exploratory Discussion",
            "desc": "A brief meeting with Smart Village leadership to share perspectives, understand current operational priorities, and identify candidate areas of synergy."
        },
        {
            "step": "STEP 2",
            "title": "Targeted Pilot Project",
            "desc": "Explore a limited, low-risk proof-of-concept in a designated visitor or event zone to empirically evaluate performance and tenant reception."
        },
        {
            "step": "STEP 3",
            "title": "Collaborative Terms",
            "desc": "Jointly structure commercial terms, operational boundaries, and mutual milestones designed specifically to match Smart Village's governance model."
        }
    ]

    step_w = Inches(3.68)
    step_h = Inches(2.2)
    step_y = Inches(2.25)

    for i, st in enumerate(s6_steps):
        x = Inches(0.8) + i * (step_w + card_gap)
        s_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, step_y, step_w, step_h)
        s_box.fill.solid()
        s_box.fill.fore_color.rgb = COLOR_BG_CARD
        s_box.line.color.rgb = COLOR_BORDER if i == 1 else COLOR_BORDER_MUTED
        s_box.line.width = Pt(1.5 if i == 1 else 1)

        tf_s = s_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = Inches(0.3)
        tf_s.margin_right = Inches(0.3)
        tf_s.margin_top = Inches(0.2)

        p_step = tf_s.paragraphs[0]
        p_step.text = st["step"]
        p_step.font.size = Pt(9)
        p_step.font.bold = True
        p_step.font.color.rgb = COLOR_CYAN

        p_st = tf_s.add_paragraph()
        p_st.text = st["title"]
        p_st.font.size = Pt(15)
        p_st.font.bold = True
        p_st.font.color.rgb = COLOR_WHITE

        p_sd = tf_s.add_paragraph()
        p_sd.text = f"\n{st['desc']}"
        p_sd.font.size = Pt(11)
        p_sd.font.color.rgb = COLOR_TEXT_BODY

    # Closing Banner with Prominent Closing Message
    closing_banner = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.65), Inches(11.733), Inches(1.95))
    closing_banner.fill.solid()
    closing_banner.fill.fore_color.rgb = COLOR_BG_CARD
    closing_banner.line.color.rgb = COLOR_BORDER
    closing_banner.line.width = Pt(1.5)

    tf_close = closing_banner.text_frame
    tf_close.word_wrap = True
    tf_close.margin_left = Inches(0.4)
    tf_close.margin_right = Inches(0.4)
    tf_close.margin_top = Inches(0.25)

    # Core Closing Quote
    p_q = tf_close.paragraphs[0]
    p_q.text = "\"Together, we can unlock new value from existing parking infrastructure.\""
    p_q.font.size = Pt(17)
    p_q.font.bold = True
    p_q.font.color.rgb = COLOR_CYAN
    p_q.alignment = PP_ALIGN.CENTER

    p_div = tf_close.add_paragraph()
    p_div.text = "────────────────────────────────────────────────────────────────────────"
    p_div.font.size = Pt(8.5)
    p_div.font.color.rgb = COLOR_BORDER_MUTED
    p_div.alignment = PP_ALIGN.CENTER

    p_contact = tf_close.add_paragraph()
    p_contact.text = "Park & Go  •  Founded by Loay Khaled  •  Cairo & Alexandria, Egypt\nEmail: Loay_khaled6@yahoo.com  |  Phone: +20 115 150 7298  |  Web: parkandgo-woad.vercel.app"
    p_contact.font.size = Pt(11.5)
    p_contact.font.color.rgb = COLOR_WHITE
    p_contact.alignment = PP_ALIGN.CENTER

    add_footer(s6, 6)

    # Save Presentation
    output_path = os.path.join(
        r"c:\Users\Loay khaled\Desktop\mobile app final al final\park-and-go-website",
        "Park_and_Go_Smart_Village_Proposal.pptx"
    )
    prs.save(output_path)
    print(f"Presentation generated successfully at: {output_path}")

if __name__ == "__main__":
    create_presentation()
